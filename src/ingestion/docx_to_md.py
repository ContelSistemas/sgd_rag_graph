#!/usr/bin/env python3
"""Conversor de documentos Word (.docx) a Markdown con frontmatter YAML para el SGD."""
import sys
from pathlib import Path
from docx import Document
import yaml
from datetime import datetime


def docx_to_markdown(docx_path: str, md_path: str = None, frontmatter: dict = None) -> str:
    """Convierte un archivo .docx a Markdown con frontmatter YAML."""
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {docx_path}")

    doc = Document(docx_path)

    md_lines = []

    # Extraer título del primer encabezado o usar el nombre del archivo
    title = docx_path.stem
    for para in doc.paragraphs:
        if para.style.name == 'Title' or (para.style.name.startswith('Heading') and para.style.name == 'Heading 0'):
            title = para.text.strip()
            break

    # Construir frontmatter por defecto
    default_frontmatter = {
        "id": docx_path.stem.lower().replace(' ', '_'),
        "title": title,
        "categoria": "importado",
        "autor": "Sistema SGD",
        "creado": datetime.now().strftime("%Y-%m-%d"),
        "tags": ["#importado", "#word"],
        "status": "activo",
        "Link to": None,
    }

    if frontmatter:
        default_frontmatter.update(frontmatter)

    md_lines.append("---")
    md_lines.append(yaml.dump(default_frontmatter, allow_unicode=True, default_flow_style=False).rstrip())
    md_lines.append("---")
    md_lines.append("")

    # Convertir párrafos y encabezados
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name

        if style_name == 'Title':
            md_lines.append(f"# {text}")
        elif style_name == 'Heading 1':
            md_lines.append(f"# {text}")
        elif style_name == 'Heading 2':
            md_lines.append(f"## {text}")
        elif style_name == 'Heading 3':
            md_lines.append(f"### {text}")
        elif style_name == 'Heading 4':
            md_lines.append(f"#### {text}")
        elif style_name == 'Heading 5':
            md_lines.append(f"##### {text}")
        elif style_name == 'Heading 6':
            md_lines.append(f"###### {text}")
        else:
            # Párrafo normal - verificar si tiene runs con negrita
            if para.runs and all(run.bold for run in para.runs if run.text.strip()):
                md_lines.append(f"**{text}**")
            else:
                md_lines.append(text)

        md_lines.append("")

    md_content = "\n".join(md_lines)

    if md_path:
        md_path = Path(md_path)
        md_path.parent.mkdir(parents=True, exist_ok=True)
        md_path.write_text(md_content, encoding="utf-8")
        print(f"Markdown guardado en: {md_path}")

    return md_content


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python docx_to_md.py <archivo.docx> [archivo.md]")
        sys.exit(1)

    docx_file = sys.argv[1]
    md_file = sys.argv[2] if len(sys.argv) > 2 else None
    docx_to_markdown(docx_file, md_file)
